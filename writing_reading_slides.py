import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation_from_template():
    template_filename = "UniWue_Template.pptx"
    
    if not os.path.exists(template_filename):
        print(f"Error: Could not find '{template_filename}' in current directory.")
        return

    prs = Presentation(template_filename)
    layouts = prs.slide_layouts

    def attach_notes(slide, notes_text):
        slide.notes_slide.notes_text_frame.text = notes_text

    def set_slide_title(slide, title_text):
        if slide.shapes.title:
            slide.shapes.title.text = title_text

    def populate_body(slide, content_items):
        body_tf = None
        for shape in slide.placeholders:
            if shape != slide.shapes.title and shape.has_text_frame and shape.placeholder_format.type.name == 'BODY':
                body_tf = shape.text_frame
                break
        
        if not body_tf:
            for shape in slide.placeholders:
                if shape != slide.shapes.title and shape.has_text_frame:
                    body_tf = shape.text_frame
                    break

        if not body_tf:
            box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.0), Inches(4.5))
            body_tf = box.text_frame
            body_tf.word_wrap = True

        body_tf.clear()

        for i, item in enumerate(content_items):
            p = body_tf.paragraphs[0] if i == 0 else body_tf.add_paragraph()
            p.space_after = Pt(10)
            
            if isinstance(item, tuple):
                head, desc = item
                r1 = p.add_run()
                r1.text = f"• {head}: "
                r1.font.bold = True
                r2 = p.add_run()
                r2.text = desc
            else:
                p.text = f"• {item}"

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Deckblatt 1 - Index 0)
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(layouts[0])
    set_slide_title(slide1, "Academic Reading & Writing")
    for ph in slide1.placeholders:
        if ph.placeholder_format.type.name == 'SUBTITLE':
            ph.text = "A Practical Guide for Seminar Papers and Theses"
            break

    attach_notes(
        slide1, 
        "Welcome slide. Introduce the session's objective: giving students a clear, "
        "practical roadmap for reading, writing, and organizing their research papers."
    )

    # -------------------------------------------------------------
    # SLIDE 2: Core Philosophy & Roadmap (Inhaltsfolie 1 - Index 6)
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(layouts[6])
    set_slide_title(slide2, "Core Message & Roadmap")
    populate_body(slide2, [
        ("The Paper is Your Proxy", "It is the primary communication tool; clear writing reflects clear thinking."),
        ("Behind-the-Scenes Work", "Coding, data collection, and surveys only become visible through the final text."),
        ("Plan Ahead", "Review the full timeline and submission requirements early to avoid last-minute panic.")
    ])

    attach_notes(
        slide2,
        """Welcome, everyone! Today, we want to give you a guide to academic reading and writing for your upcoming seminar papers and theses.

Whether you are working on an empirical project with data analysis or conducting your own primary data collection, the central principle for this semester is simple: Your paper is the tool to let the reader understand your research, whatever it is.

While you will spend a lot of time working behind the scenes—coding, gathering data, or designing surveys—the reader only sees the final text. A clear and well-structured paper is what makes all your hard work and findings accessible.

To make this process as smooth as possible, we strongly encourage you to take a look at the entire roadmap and review all provided documents right at the start.

Having a complete overview of the timeline, submission requirements, and guidelines from day one helps you plan ahead and avoid last-minute stress."""
    )

    # -------------------------------------------------------------
    # SLIDE 3: Literature Reading & Tracking (Inhaltsfolie 1 - Index 6)
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(layouts[6])
    set_slide_title(slide3, "Literature Reading & Transparent Tracking")
    populate_body(slide3, [
        ("Filter Literature", "Check title, abstract, and conclusion before reading a paper in full."),
        ("The Relevance Test", "Ensure every cited paper adds concrete value to your project."),
        ("Methodological Baseline", "Use established literature to guide your empirical strategy and design choices."),
        ("Direct Reading & Organization", "Read relevant papers thoroughly yourself and maintain tracking tables and annotated PDFs.")
    ])

    attach_notes(
        slide3,
        """When it comes to reading literature, the goal is to work smartly rather than trying to read every single paper from start to finish like a novel. Start broadly to understand the general landscape of your topic, but quickly focus in on the literature that is directly useful for your specific research.

A great first step is to check the title, abstract, and the introduction or conclusion first. If those sections don't align with your topic or research question, you don't need to read the entire paper.

Keep a clear relevance test in mind: if a paper is too general and you cannot identify a specific sentence, paragraph, or clear reason to include it, it's usually best not to cite it. Always ask yourself what concrete value a paper adds to your project.

Keep in mind that literature isn't just there for theoretical background—it also provides a great baseline for your methodology. Looking at how established papers construct their empirical strategy, design their surveys, or structure their questions gives you a solid foundation for your own choices.

Make sure to read the relevant papers thoroughly yourself rather than relying on automated summaries. To stay organized and save time later: keep a simple literature table noting why each paper is relevant, save PDFs with clear notes, and use reference management software from the beginning."""
    )

    # -------------------------------------------------------------
    # SLIDE 4: Visual Elements / Conciseness (Inhaltsfolie 3 - Index 8)
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(layouts[8])
    set_slide_title(slide4, "Conciseness & Visual Elements")
    populate_body(slide4, [
        ("Prioritize Visuals", "Use tables and figures for conciseness and clear presentation of data."),
        ("Avoid Clutter", "Ensure every element adds direct value to your main narrative.")
    ])

    box_visual = slide4.shapes.add_textbox(Inches(6.8), Inches(1.8), Inches(5.5), Inches(4.8))
    tf_v = box_visual.text_frame
    tf_v.word_wrap = True
    p_v = tf_v.paragraphs[0]
    p_v.text = "[ VISUAL PLACEHOLDER ]\n\nInsert example table, figure, or infogram here."
    p_v.font.size = Pt(14)
    p_v.font.bold = True

    attach_notes(
        slide4,
        "Emphasize to students that good academic writing prioritizes conciseness, making strong use of visual elements like tables and figures to keep information clean and direct."
    )

    # -------------------------------------------------------------
    # SLIDE 5: Structuring & Writing Process (Inhaltsfolie 5 - Index 10)
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(layouts[10])
    set_slide_title(slide5, "Structuring and Writing the Paper")
    populate_body(slide5, [
        ("Cyclical Process", "Writing is a cycle of drafting, reading, correcting, rewriting, and reading again—start early."),
        ("Draft Sharing", "We may ask for paper drafts; we don't expect perfection, just share what you have on paper."),
        ("Frontload Arguments", "State your research question and core contribution clearly near the beginning.")
    ])

    attach_notes(
        slide5,
        """When it comes to the writing process itself, the key is to realize that your project and your paper develop together. You don't need to wait until all your analysis or data collection is completely finished before you start putting words on paper.

It helps to remember that writing is not a straight line, but rather a cycle where you write, read, correct, rewrite, and read again. Because of this cyclical nature, starting early is essential.

Along the way, we might ask for some paper drafts for some of you. We do not expect you to write a polished draft specifically for us at that point—we just want you to share what you actually have on paper up to that point.

As your draft evolves, frontload your main argument and state your research question early on so the reader knows right away where the paper is heading."""
    )

    # -------------------------------------------------------------
    # SLIDE 6: University Support & Next Steps (Inhaltsfolie 1 - Index 6)
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(layouts[6])
    set_slide_title(slide6, "University Support & Next Steps")
    populate_body(slide6, [
        ("Institutional Resources", "Utilize the university library and Writing Center (Schreibzentrum) for support."),
        ("Paper Guidelines", "Refer to the official paper guidelines document for further details.")
    ])

    attach_notes(
        slide6,
        """Finally, remember that learning to navigate academic reading and writing is a skill you build over time, and you don't have to figure it all out completely alone.

Throughout your project, make use of the institutional support available across the university. Services like the university library and the Writing Center (Schreibzentrum) offer valuable workshops and consultations.

You will find further details in our official Paper Guidelines document. Good luck with your projects—we are looking forward to seeing your work develop!"""
    )

    # -------------------------------------------------------------
    # SLIDE 7: Closing Slide (Abschlussseite - Index 20)
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(layouts[20])
    set_slide_title(slide7, "Thank You & Good Luck!")

    attach_notes(slide7, "Final slide. Wrap up the session.")

    output_filename = "academic_reading_and_writing_presentation.pptx"
    prs.save(output_filename)
    print(f"\nPresentation successfully generated and saved as '{output_filename}'!")

if __name__ == "__main__":
    create_presentation_from_template()